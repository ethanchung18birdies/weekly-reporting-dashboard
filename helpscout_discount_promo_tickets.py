#!/usr/bin/env python3
"""Export the newest Help Scout tickets mentioning discount or promo terms."""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from datetime import datetime, timedelta, timezone
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode, urljoin
from urllib.request import Request, urlopen

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError as exc:
    raise SystemExit(
        "This script needs openpyxl. Run it with your Help Scout environment:\n"
        "~/Desktop/helpscout-venv/bin/python ~/Desktop/helpscout_discount_promo_tickets.py"
    ) from exc


API_BASE = "https://api.helpscout.net/v2"
TOKEN_URL = f"{API_BASE}/oauth2/token"
SEARCH_TERMS = ("discount code", "discount link", "promo")
DEFAULT_LIMIT = 100
DEFAULT_OUTPUT = Path.home() / "Desktop" / "helpscout_discount_promo_tickets_last_year.xlsx"
MAX_RETRIES = 6


class HTMLTextExtractor(HTMLParser):
    BLOCK_TAGS = {
        "address", "article", "aside", "blockquote", "br", "div", "dl", "dt",
        "dd", "fieldset", "figcaption", "figure", "footer", "form", "h1", "h2",
        "h3", "h4", "h5", "h6", "header", "hr", "li", "main", "nav", "ol",
        "p", "pre", "section", "table", "tr", "td", "th", "thead", "tbody",
        "tfoot", "ul",
    }

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.skip_depth = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in {"script", "style"}:
            self.skip_depth += 1
        elif self.skip_depth == 0 and tag in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self.skip_depth:
            self.skip_depth -= 1
        elif self.skip_depth == 0 and tag in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self.skip_depth == 0 and data:
            self.parts.append(data)

    def text(self) -> str:
        value = unescape("".join(self.parts)).replace("\r", "\n")
        value = re.sub(r"[ \t]+\n", "\n", value)
        value = re.sub(r"\n[ \t]+", "\n", value)
        value = re.sub(r"[ \t]+", " ", value)
        return re.sub(r"\n{3,}", "\n\n", value).strip()


def strip_html(value: str) -> str:
    parser = HTMLTextExtractor()
    parser.feed(value or "")
    parser.close()
    return parser.text()


def parse_created_at(value: object) -> datetime | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def excel_date(value: object) -> datetime | str:
    parsed = parse_created_at(value)
    return parsed.replace(tzinfo=None) if parsed else ""


def conversation_url(conversation: dict) -> str:
    href = ((conversation.get("_links") or {}).get("web") or {}).get("href")
    if href:
        return str(href)
    conversation_id, number = conversation.get("id"), conversation.get("number")
    if conversation_id and number:
        return f"https://secure.helpscout.net/conversation/{conversation_id}/{number}/"
    return ""


def thread_is_customer(thread: dict) -> bool:
    creator = thread.get("createdBy") or {}
    return str(creator.get("type") or "").lower() == "customer" or bool(thread.get("customer"))


def collect_message_text(conversation: dict) -> tuple[str, str]:
    customer_messages: list[str] = []
    agent_messages: list[str] = []
    threads = conversation.get("_embedded", {}).get("threads", []) or []
    for thread in sorted(threads, key=lambda item: item.get("createdAt") or ""):
        if str(thread.get("state") or "").lower() not in {"", "published"}:
            continue
        if str(thread.get("type") or "").lower() in {"note", "lineitem", "chatline"}:
            continue
        message = strip_html(
            str(thread.get("body") or thread.get("plaintext") or thread.get("bodyPreview") or "")
        )
        if message:
            (customer_messages if thread_is_customer(thread) else agent_messages).append(message)
    return "\n\n---\n\n".join(customer_messages), "\n\n---\n\n".join(agent_messages)


def render_progress(label: str, current: int, total: int) -> None:
    width = 28
    filled = int(width * current / max(total, 1))
    bar = "#" * filled + "-" * (width - filled)
    print(f"\r{label} [{bar}] {current}/{total}", end="", file=sys.stderr, flush=True)


class HelpScoutClient:
    def __init__(self, client_id: str, client_secret: str, pause_seconds: float) -> None:
        self.client_id = client_id
        self.client_secret = client_secret
        self.pause_seconds = pause_seconds
        self.access_token: str | None = None

    def authenticate(self) -> None:
        payload = urlencode(
            {
                "grant_type": "client_credentials",
                "client_id": self.client_id,
                "client_secret": self.client_secret,
            }
        ).encode("utf-8")
        request = Request(
            TOKEN_URL,
            data=payload,
            headers={"Content-Type": "application/x-www-form-urlencoded", "Accept": "application/json"},
            method="POST",
        )
        try:
            with urlopen(request, timeout=30) as response:
                self.access_token = json.loads(response.read().decode("utf-8")).get("access_token")
        except (HTTPError, URLError) as exc:
            raise RuntimeError(f"Could not sign in to Help Scout: {exc}") from exc
        if not self.access_token:
            raise RuntimeError("Help Scout did not return an access token.")

    def request_json(self, path: str, **params) -> dict:
        if not self.access_token:
            raise RuntimeError("The Help Scout client is not authenticated.")
        url = urljoin(f"{API_BASE}/", path.lstrip("/"))
        query = urlencode({key: value for key, value in params.items() if value is not None})
        if query:
            url = f"{url}?{query}"
        request = Request(url, headers={"Authorization": f"Bearer {self.access_token}", "Accept": "application/json"})

        for attempt in range(1, MAX_RETRIES + 1):
            try:
                with urlopen(request, timeout=60) as response:
                    payload = json.loads(response.read().decode("utf-8"))
                if self.pause_seconds:
                    time.sleep(self.pause_seconds)
                return payload
            except HTTPError as exc:
                if exc.code == 404:
                    return {}
                if exc.code not in {429, 500, 502, 503, 504} or attempt == MAX_RETRIES:
                    body = exc.read().decode("utf-8", errors="replace")
                    raise RuntimeError(f"Help Scout API error {exc.code} for {url}: {body}") from exc
                reason = f"HTTP {exc.code}"
            except (URLError, ConnectionResetError, TimeoutError) as exc:
                if attempt == MAX_RETRIES:
                    raise RuntimeError(f"Could not reach Help Scout API at {url}: {exc}") from exc
                reason = "network error"

            wait_seconds = min(30, 2**attempt)
            print(f"\nRetrying after {reason} in {wait_seconds}s...", file=sys.stderr)
            time.sleep(wait_seconds)
        raise RuntimeError(f"Exceeded retry limit for {url}")

    def search_term(self, term: str, cutoff: datetime, per_term_limit: int) -> list[dict]:
        results: list[dict] = []
        page = 1
        while len(results) < per_term_limit:
            payload = self.request_json(
                "/conversations",
                status="all",
                query=f'"{term}"',
                sortField="createdAt",
                sortOrder="desc",
                page=page,
                pageSize=100,
            )
            conversations = payload.get("_embedded", {}).get("conversations", []) or []
            for conversation in conversations:
                created_at = parse_created_at(conversation.get("createdAt"))
                if created_at and created_at >= cutoff:
                    results.append(conversation)

            page_info = payload.get("page", {}) or {}
            total_pages = int(page_info.get("totalPages") or page_info.get("pages") or page)
            render_progress(f"Finding {term}", page, total_pages)
            oldest = parse_created_at(conversations[-1].get("createdAt")) if conversations else None
            if page >= total_pages or not conversations or (oldest and oldest < cutoff):
                print(file=sys.stderr)
                return results[:per_term_limit]
            page += 1
        print(file=sys.stderr)
        return results[:per_term_limit]

    def get_conversation(self, conversation_id: int) -> dict:
        return self.request_json(f"/conversations/{conversation_id}", embed="threads")


def write_workbook(rows: list[dict], output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Discount and Promo Tickets"
    headers = [
        "Ticket ID",
        "Help Scout URL",
        "Date Created",
        "Status",
        "Matched Term(s)",
        "Customer Feedback",
        "Agent Response",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="12372A")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in rows:
        sheet.append(
            [
                row["ticket_id"],
                row["url"],
                row["created_at"],
                row["status"],
                row["matched_terms"],
                row["customer_feedback"],
                row["agent_response"],
            ]
        )
        row_number = sheet.max_row
        url_cell = sheet.cell(row_number, 2)
        if row["url"]:
            url_cell.hyperlink = row["url"]
            url_cell.style = "Hyperlink"
        sheet.cell(row_number, 3).number_format = "mm/dd/yy h:mm AM/PM"
        for column in range(1, len(headers) + 1):
            sheet.cell(row_number, column).alignment = Alignment(wrap_text=True, vertical="top")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G{max(sheet.max_row, 1)}"
    for letter, width in {"A": 14, "B": 58, "C": 22, "D": 14, "E": 30, "F": 90, "G": 90}.items():
        sheet.column_dimensions[letter].width = width
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def export_tickets(client: HelpScoutClient, cutoff: datetime, limit: int, output_path: Path) -> int:
    summaries: dict[str, dict] = {}
    matched_terms: dict[str, set[str]] = {}
    for term in SEARCH_TERMS:
        # One hundred per term is enough to determine the newest 100 overall matches.
        for summary in client.search_term(term, cutoff, limit):
            conversation_id = str(summary.get("id") or "")
            if not conversation_id:
                continue
            summaries[conversation_id] = summary
            matched_terms.setdefault(conversation_id, set()).add(term)

    ordered = sorted(
        summaries.values(),
        key=lambda item: parse_created_at(item.get("createdAt")) or datetime.min.replace(tzinfo=timezone.utc),
        reverse=True,
    )[:limit]
    rows: list[dict] = []
    for index, summary in enumerate(ordered, start=1):
        conversation_id = str(summary.get("id") or "")
        conversation = client.get_conversation(int(conversation_id)) if conversation_id else {}
        source = conversation or summary
        customer_feedback, agent_response = collect_message_text(conversation) if conversation else ("", "")
        rows.append(
            {
                "ticket_id": str(source.get("number") or source.get("id") or ""),
                "url": conversation_url(source),
                "created_at": excel_date(source.get("createdAt")),
                "status": str(source.get("status") or ""),
                "matched_terms": " | ".join(sorted(matched_terms.get(conversation_id, set()))),
                "customer_feedback": customer_feedback,
                "agent_response": agent_response,
            }
        )
        render_progress("Fetching ticket details", index, len(ordered))
    if ordered:
        print(file=sys.stderr)

    write_workbook(rows, output_path)
    return len(rows)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Export up to 100 past-year Help Scout tickets mentioning discount code, discount link, or promo."
    )
    parser.add_argument("--limit", type=int, default=DEFAULT_LIMIT, help="Maximum unique tickets to export (default: 100).")
    parser.add_argument("--days", type=int, default=365, help="Rolling lookback period in days (default: 365).")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help=f"Workbook to create (default: {DEFAULT_OUTPUT})")
    parser.add_argument("--pause-seconds", type=float, default=0.1, help="Delay between Help Scout requests (default: 0.1).")
    args = parser.parse_args()
    if args.limit < 1 or args.days < 1:
        print("--limit and --days must both be at least 1.", file=sys.stderr)
        return 2

    client_id = os.environ.get("HELPSCOUT_CLIENT_ID") or os.environ.get("HELPSCOUT_APP_ID")
    client_secret = os.environ.get("HELPSCOUT_CLIENT_SECRET") or os.environ.get("HELPSCOUT_APP_SECRET")
    if not client_id or not client_secret:
        print(
            "Missing Help Scout credentials. Set HELPSCOUT_CLIENT_ID and "
            "HELPSCOUT_CLIENT_SECRET (or HELPSCOUT_APP_ID and HELPSCOUT_APP_SECRET).",
            file=sys.stderr,
        )
        return 2

    client = HelpScoutClient(client_id, client_secret, args.pause_seconds)
    client.authenticate()
    output_path = args.output.expanduser().resolve()
    cutoff = datetime.now(timezone.utc) - timedelta(days=args.days)
    count = export_tickets(client, cutoff, args.limit, output_path)
    print(f"Wrote {count} tickets to {output_path}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
